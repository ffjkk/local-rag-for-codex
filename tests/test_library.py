import io,json,time,sys,zipfile,asyncio,os
from pathlib import Path
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import knowledge as kb
from parsers import parse
import app
from fastapi.testclient import TestClient
HEADERS={'X-RAG-UI':'1'}
@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(kb,'DB',tmp_path/'test.sqlite3'); kb.init()
    with TestClient(app.app) as c: yield c

def create(folder,name,text): return kb.import_document(folder,name,text.encode(),text)

def test_crud_and_index_atomic(client):
    assert client.post('/api/folders',json={'path':'项目/服务器'},headers=HEADERS).status_code==200
    doc=create('项目/服务器','setup.md','七日杀获取服务器 ID 失败，把 steamnetworking 改为空白。')
    assert kb.search('服务器ID失败',scope='all')['results'][0]['document_id']==doc['id']
    changed=client.put('/api/documents/'+doc['id'],json={'content':'唯一更新词 STONEBRIDGE：新配置端口是 23456。','version':1},headers=HEADERS)
    assert changed.status_code==200 and changed.json()['version']==2
    result=kb.search('STONEBRIDGE',scope='all')['results']
    assert result and result[0]['version']==2 and '23456' in result[0]['content']
    assert all('steamnetworking' not in x['content'] for x in kb.search('steamnetworking',scope='all')['results'])
    assert client.put('/api/documents/'+doc['id'],json={'content':'stale','version':1},headers=HEADERS).status_code==400
    assert kb.get_document(doc['id'])['version']==2
    assert client.delete('/api/documents/'+doc['id'],headers=HEADERS).status_code==200
    assert not kb.search('STONEBRIDGE',scope='all')['results']

def test_category_routing_and_all(client):
    create('游戏服务器','minecraft.md','Minecraft 游戏服务器支持 PaperMC 插件，需要配置 Java 运行环境。')
    create('烘焙食谱','蛋糕.txt','巧克力蛋糕需要鸡蛋、面粉、牛奶和黄油，放入烤箱烘焙。')
    create('旅行/日本','京都.md','京都旅行可以乘坐电车参观清水寺，春天观赏樱花。')
    auto=kb.search('怎么做巧克力蛋糕，需要什么材料')
    assert '烘焙食谱' in auto['selected_categories']
    assert auto['results'][0]['category']=='烘焙食谱'
    assert len(auto['selected_categories'])<3
    all_result=kb.search('资料',scope='all')
    assert len(all_result['selected_categories'])==3
    assert kb.search('电车',category='旅行')['selected_categories']==['旅行/日本']
    assert not kb.search('电车',category='不存在')['results']

def test_folders_and_filter(client):
    create('A/B','one.txt','绿色文档内容')
    assert client.delete('/api/folders?path=A',headers=HEADERS).status_code==400
    assert client.get('/api/list?q=绿色').json()['documents'][0]['name']=='one.txt'
    assert client.delete('/api/folders?path=A&recursive=true',headers=HEADERS).json()['deleted_documents']==1
    assert kb.status()['documents']==0 and kb.status()['chunks']==0
    assert client.delete('/api/folders?path=',headers=HEADERS).status_code==400
    assert client.get('/api/list?folder=A').status_code==404

@pytest.mark.parametrize('bad',['../outside','A/../../B','/abs','C:\\bad','a//b','a/./b'])
def test_path_validation(client,bad):
    assert client.post('/api/folders',json={'path':bad},headers=HEADERS).status_code==400

def wait_job(client,job):
    for _ in range(400):
        status=client.get('/api/jobs/'+job).json()
        if status['state']=='done': return status
        time.sleep(.05)
    pytest.fail('upload timeout')

def test_batch_upload_tree_and_errors(client):
    files=[('files',('a.md','# 名称\n这是批量上传文档'.encode(),'text/plain')),('files',('b.txt','嵌套目录资料'.encode(),'text/plain')),('files',('bad.exe',b'MZ\x00\x01','application/octet-stream'))]
    r=client.post('/api/upload',files=files,data={'paths':json.dumps(['上传目录/a.md','上传目录/子目录/b.txt','bad.exe']),'folder':''},headers=HEADERS)
    assert r.status_code==200,r.text
    status=wait_job(client,r.json()['job_id'])
    assert [x['status'] for x in status['items']]==['success','success','error']
    assert kb.listing('上传目录/子目录')['documents'][0]['name']=='b.txt'
    duplicate=client.post('/api/upload',files=[files[0]],data={'paths':json.dumps(['上传目录/a.md'])},headers=HEADERS)
    assert wait_job(client,duplicate.json()['job_id'])['items'][0]['status']=='error'
    assert kb.status()['documents']==2

def test_security(client):
    assert client.post('/api/folders',json={'path':'bad'}).status_code==403
    assert client.get('/api/status',headers={'origin':'https://evil.example'}).status_code==403
    assert client.get('/api/status',headers={'host':'evil.example'}).status_code==400
    r=client.post('/api/upload',files={'files':('a.txt',b'x')},data={'paths':'["../bad.txt"]'},headers=HEADERS)
    assert r.status_code==400

def test_parsers():
    assert parse('a.txt','中文测试'.encode('gb18030'))[0]=='中文测试'
    from docx import Document
    d=Document();d.add_paragraph('Word 段落内容');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='列一';t.cell(0,1).text='列二';b=io.BytesIO();d.save(b)
    for name in ('a.docx','a.docs'):
        text,_=parse(name,b.getvalue());assert 'Word 段落内容' in text and '列一 | 列二' in text
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    w=PdfWriter();p=w.add_blank_page(width=400,height=400)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    p[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 18 Tf 20 350 Td (PDF Knowledge parser verification text) Tj ET');p[NameObject('/Contents')]=w._add_object(stream)
    b=io.BytesIO();w.write(b);text,_=parse('sample.pdf',b.getvalue());assert 'PDF Knowledge' in text and '第 1 页' in text

def test_scanned_pdf_ocr():
    from PIL import Image,ImageDraw,ImageFont
    image=Image.new('RGB',(1000,350),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48)
    draw.text((40,120),'LOCAL KNOWLEDGE TEST',fill='black',font=font);b=io.BytesIO();image.save(b,format='PDF')
    text,warning=parse('scan.pdf',b.getvalue());assert 'KNOWLEDGE' in text.upper() and 'OCR' in warning


def test_unrelated_query_does_not_force_local_answer(client):
    create('游戏服务器','minecraft.md','Minecraft 游戏服务器支持 PaperMC 插件，需要配置 Java 运行环境。')
    assert not kb.search('2加2等于多少')['results']
    assert not kb.search('今天北京天气如何')['results']


@pytest.mark.parametrize('name',['main.js','worker.py','types.ts','view.vue','query.sql','settings.json','config.yaml','table.csv','Dockerfile','.gitignore','notes.custom'])
def test_generic_text_formats(name):
    source='# 中文说明\nprint("hello")\n'
    assert parse(name,source.encode('utf-8'))[0]==source

@pytest.mark.parametrize('encoding',['utf-8-sig','utf-16','utf-32','gb18030'])
def test_text_encodings(encoding):
    source='中文注释\n变量 = 42\n'
    assert parse('code.py',source.encode(encoding))[0]==source

@pytest.mark.parametrize('raw',[b'PK\x03\x04fake',b'MZbinary',b'\x89PNG\r\n',b'\x7fELF',b'abc\x00def',b'abc\x01def',b'\xff'])
def test_binary_content_rejected_even_with_code_extension(raw):
    with pytest.raises(ValueError): parse('renamed.py',raw)

def test_code_upload_edit_and_search(client):
    source='def KNOWLEDGE_CODE_MARKER():\n    return "代码入库测试"\n'
    r=client.post('/api/upload',files=[('files',('example.py',source.encode(),'text/x-python'))],data={'paths':'["源码/example.py"]'},headers=HEADERS)
    result=wait_job(client,r.json()['job_id'])['items'][0]
    assert result['status']=='success'
    doc=kb.get_document(result['id']); assert doc['content']==source
    hits=kb.search('KNOWLEDGE_CODE_MARKER',scope='all')['results']
    assert hits and hits[0]['document_id']==result['id']
    assert client.put('/api/documents/'+result['id'],json={'content':'# UPDATED_CODE_MARKER\nvalue = 23456','version':1},headers=HEADERS).status_code==200
    assert '23456' in kb.search('UPDATED_CODE_MARKER',scope='all')['results'][0]['content']


def test_unclassified_always_in_retrieval(client):
    create('未分类','shared.md','默认知识 BASELINE_MARKER；所有项目都需要参照。')
    create('未分类/补充','extra.md','补充通用知识 EXTRA_MARKER。')
    create('游戏','server.md','游戏服务器端口配置。')
    for options in [{},{'category':'游戏'},{'category':'不存在'},{'scope':'all'}]:
        result=kb.search('BASELINE_MARKER',**options)
        assert {'未分类','未分类/补充'}<=set(result['selected_categories'])
        assert set(result['always_searched_categories'])=={'未分类','未分类/补充'}
        assert len(result['selected_categories'])==len(set(result['selected_categories']))
        assert any(x['filename']=='shared.md' for x in result['results'])
    assert client.delete('/api/folders?path=未分类&recursive=true',headers=HEADERS).status_code==400
    assert client.post('/api/folders/move',json={'source':'未分类','target_parent':'游戏'},headers=HEADERS).status_code==400


def test_move_document_and_conflicts(client):
    kb.create_folder('目的地')
    d=create('原目录','source.py','def MOVE_MARKER():\n    return 23456')
    result=client.post('/api/documents/'+d['id']+'/move',json={'target_folder':'目的地','version':1},headers=HEADERS)
    assert result.status_code==200,result.text
    moved=result.json();assert moved['folder']=='目的地' and moved['version']==2
    assert kb.get_document(d['id'])['content']==moved['content']
    with kb.connect() as c:
        assert not c.execute("SELECT 1 FROM routes WHERE folder='原目录'").fetchone()
        assert c.execute('PRAGMA foreign_key_check').fetchall()==[]
    hits=kb.search('MOVE_MARKER',category='目的地')['results']
    assert hits[0]['document_id']==d['id'] and hits[0]['category']=='目的地'
    assert client.post('/api/documents/'+d['id']+'/move',json={'target_folder':'原目录','version':1},headers=HEADERS).status_code==400
    create('原目录','source.py','另一个文件')
    assert client.post('/api/documents/'+d['id']+'/move',json={'target_folder':'原目录','version':2},headers=HEADERS).status_code==400
    assert kb.get_document(d['id'])['folder']=='目的地'


def test_move_folder_subtree_and_cycle(client):
    doc=create('源/子目录','tree.md','SUBTREE_MARKER 子目录知识')
    kb.create_folder('源/空目录');kb.create_folder('目标')
    assert client.post('/api/folders/move',json={'source':'源','target_parent':'源/子目录'},headers=HEADERS).status_code==400
    r=client.post('/api/folders/move',json={'source':'源','target_parent':'目标'},headers=HEADERS)
    assert r.status_code==200,r.text
    assert kb.get_document(doc['id'])['folder']=='目标/源/子目录'
    assert kb.listing('目标/源/空目录')['documents']==[]
    assert kb.search('SUBTREE_MARKER',category='目标')['results'][0]['category']=='目标/源/子目录'
    with kb.connect() as c:
        assert not c.execute("SELECT 1 FROM folders WHERE path='源'").fetchone()
        assert c.execute('PRAGMA foreign_key_check').fetchall()==[]
    kb.create_folder('源')
    assert client.post('/api/folders/move',json={'source':'目标/源','target_parent':''},headers=HEADERS).status_code==400
    assert kb.get_document(doc['id'])['folder']=='目标/源/子目录'


def test_failed_move_rolls_back(client,monkeypatch):
    d=create('原/子','a.md','事务回滚测试')
    kb.create_folder('新')
    def broken(*args,**kwargs): raise ValueError('模拟向量计算失败')
    monkeypatch.setattr(kb,'prepare',broken)
    r=client.post('/api/folders/move',json={'source':'原','target_parent':'新'},headers=HEADERS)
    assert r.status_code==400
    assert kb.get_document(d['id'])['folder']=='原/子'
    with kb.connect() as c: assert not c.execute("SELECT 1 FROM folders WHERE path='新/原'").fetchone()


def test_safe_render_modes(client):
    d=create('示例','preview.md','# 标题\n\n**加粗**\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n<script>alert(1)</script>\n\n[危险](javascript:alert(1))\n\n![图片](https://evil.example/image)')
    r=client.get('/api/documents/'+d['id']+'/render');assert r.status_code==200
    html=r.json()['html'];assert '<h1>标题</h1>' in html and '<strong>加粗</strong>' in html and '<table>' in html
    assert '<script' not in html and '<img' not in html and 'href="javascript:' not in html
    d=create('示例','sample.json','{"title":"中文","count":2}')
    assert '中文' in client.get('/api/documents/'+d['id']+'/render').json()['html']
    d=create('示例','sample.csv','列一,列二\n<script>,2')
    html=client.get('/api/documents/'+d['id']+'/render').json()['html']
    assert '<table>' in html and '<script>' not in html
